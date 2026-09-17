import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

MD_PATH = r"Docs/Buku_Manual_Pengguna_ESCM_Marmer.md"
DOCX_OUTPUT = r"Docs/laporan_kegiatan/DOCX/Buku_Manual_Pengguna_ESCM_Marmer.docx"
PNG_DIR = r"Docs/laporan_kegiatan/PNG"

# Mapping of keywords / chapters to screenshots and captions
SCREENSHOT_MAPPINGS = [
    {
        "pattern": r"### 1\.1 Latar Belakang",
        "image": "manual_01_beranda_landing_page.png",
        "caption": "Gambar 1.1: Halaman Beranda (Landing Page) Publik Klaster IKM Kerajinan Marmer Tulungagung"
    },
    {
        "pattern": r"### 1\.2 Arsitektur Sistem",
        "image": "manual_18_supply_chain_flow.png",
        "caption": "Gambar 1.2: Diagram Visualisasi Arsitektur Alur Rantai Pasok Hulu-ke-Hilir IKM Marmer"
    },
    {
        "pattern": r"### 2\.2 Hak Akses Pengguna",
        "image": "manual_08_login_admin.png",
        "caption": "Gambar 2.1: Halaman Autentikasi Pengguna & Pintu Masuk Multi-Role RBAC"
    },
    {
        "pattern": r"### 2\.3 Manajemen Akun Pengguna",
        "image": "manual_20_user_management.png",
        "caption": "Gambar 2.2: Antarmuka Manajemen Pengguna Multi-Role RBAC dan Status Akun Operator"
    },
    {
        "pattern": r"### 3\.1 Halaman Etalase & Katalog Produk",
        "image": "manual_02_katalog_multi_filter.png",
        "caption": "Gambar 3.1: Antarmuka Etalase Publik dengan Multi-Filter Kategori & Toko IKM"
    },
    {
        "pattern": r"2\. \*\*Kartu Produk E-Commerce:\*\*",
        "image": "manual_03_detail_produk.png",
        "caption": "Gambar 3.2: Halaman Rincian Spesifikasi Teknis Produk Kerajinan & Tombol Konsultasi WhatsApp"
    },
    {
        "pattern": r"### 3\.2 Alur Transaksi Checkout",
        "image": "manual_04_checkout_form.png",
        "caption": "Gambar 3.3: Formulir Checkout E-Commerce dengan Pilihan Skema DP 50% & Midtrans"
    },
    {
        "pattern": r"### 3\.3 Halaman Faktur Tagihan Digital",
        "image": "manual_05_invoice_digital.png",
        "caption": "Gambar 3.4: Faktur Tagihan Digital Interaktif Dilengkapi Snap Pay & Fitur Anti-Kehilangan Tagihan"
    },
    {
        "pattern": r"Pembayaran Online Midtrans Snap",
        "image": "manual_06_midtrans_snap.png",
        "caption": "Gambar 3.5: Popup Pembayaran Instan Midtrans Snap (QRIS Nasional, VA Bank & E-Wallet)"
    },
    {
        "pattern": r"### 3\.4 Halaman Pelacakan Pesanan Real-Time",
        "image": "manual_07_lacak_pesanan.png",
        "caption": "Gambar 3.6: Halaman Pelacakan Pesanan Real-Time 5-Tahap dan Cache Riwayat Browser"
    },
    {
        "pattern": r"### 4\.2 Halaman Manajemen Pesanan",
        "image": "manual_10_orders_management.png",
        "caption": "Gambar 4.1: Dashboard Manajemen Pesanan Admin dengan Tombol 2-Gate Verifikasi SPK"
    },
    {
        "pattern": r"### 5\.1 Halaman Master Produk",
        "image": "manual_19_master_produk.png",
        "caption": "Gambar 5.1: Antarmuka Master Produk Kerajinan Marmer & Onyx Terkurasi Multi-IKM"
    },
    {
        "pattern": r"### 6\.1 Dashboard Analitik Eksekutif",
        "image": "manual_09_dashboard_kpi.png",
        "caption": "Gambar 6.1: Dashboard Eksekutif KPI Operasional, Notifikasi Order Baru & Tren Produksi"
    },
    {
        "pattern": r"### 6\.2 Visualisasi Alur Rantai Pasok",
        "image": "manual_18_supply_chain_flow.png",
        "caption": "Gambar 6.2: Alur Komprehensif 7-Tahap Rantai Pasok IKM Marmer Campurdarat"
    },
    {
        "pattern": r"### 7\.1 Master Bahan Baku",
        "image": "manual_11_materials_stock.png",
        "caption": "Gambar 7.1: Antarmuka Manajemen Persediaan Bongkahan Batu Tambang & Batas Safety Stock"
    },
    {
        "pattern": r"### 8\.1 Papan Kanban Produksi Digital",
        "image": "manual_12_kanban_board.png",
        "caption": "Gambar 8.1: Papan Kanban Penjadwalan Produksi Digital 5 Stasiun Kerja Bengkel"
    },
    {
        "pattern": r"### 8\.4 Monitoring Barang Dalam Proses",
        "image": "manual_13_wip_tracking.png",
        "caption": "Gambar 8.2: Pelacakan Work-In-Progress (WIP) dan Utilisasi Mesin Bubut Pengrajin"
    },
    {
        "pattern": r"### 9\.1 Inspeksi Kualitas 2-Tahap",
        "image": "manual_14_qc_inspection.png",
        "caption": "Gambar 9.1: Formulir Inspeksi Mutu 2-Tahap (Struktur & Polesan) dan Riwayat Pengujian"
    },
    {
        "pattern": r"### 9\.2 Pencatatan & Pengendalian Limbah Industri",
        "image": "manual_15_waste_logs.png",
        "caption": "Gambar 9.2: Pencatatan Limbah Padat & Lumpur Slurry untuk Hilirisasi Industri Hijau"
    },
    {
        "pattern": r"### 10\.1 Manajemen Pengiriman Kargo",
        "image": "manual_16_distribution_packing.png",
        "caption": "Gambar 10.1: Antarmuka Pengiriman Kargo & Checklist Verifikasi Packing Peti Kayu Solid"
    },
    {
        "pattern": r"### 11\.1 Halaman Forecasting AI",
        "image": "manual_17_forecasting_arima.png",
        "caption": "Gambar 11.1: Grafik Peramalan Permintaan AI Menggunakan Model Terbaik ARIMA(2,0,2)"
    }
]

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_callout_box(doc, text, title="CATATAN PENTING"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "F0F9FF")
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    # Border: left thick blue, others none
    tcPr = cell._element.get_or_add_tcPr()
    tcBorders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="none"/>
            <w:left w:val="single" w:sz="36" w:space="0" w:color="0284C7"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(tcBorders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    run_title = p.add_run(f"📌 {title}: ")
    run_title.bold = True
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(10)
    run_title.font.color.rgb = RGBColor(3, 105, 161)
    
    run_text = p.add_run(text)
    run_text.font.name = "Calibri"
    run_text.font.size = Pt(9.5)
    run_text.font.color.rgb = RGBColor(15, 23, 42)
    
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(4)

def format_table(table):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_row = table.rows[0]
    for cell in header_row.cells:
        set_cell_background(cell, "1E3A8A")
        set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
        for p in cell.paragraphs:
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            for r in p.runs:
                r.bold = True
                r.font.name = "Calibri"
                r.font.size = Pt(9.5)
                r.font.color.rgb = RGBColor(255, 255, 255)
                
    for i, row in enumerate(table.rows[1:], start=1):
        bg_color = "F8FAFC" if i % 2 == 1 else "FFFFFF"
        for cell in row.cells:
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=90, bottom=90, left=140, right=140)
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                for r in p.runs:
                    r.font.name = "Calibri"
                    r.font.size = Pt(9)
                    r.font.color.rgb = RGBColor(30, 41, 59)
                    
    # Table borders
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(f'''
            <w:tblBorders {nsdecls("w")}>
                <w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
                <w:bottom w:val="single" w:sz="6" w:space="0" w:color="94A3B8"/>
                <w:insideH w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>
                <w:insideV w:val="none"/>
                <w:left w:val="none"/>
                <w:right w:val="none"/>
            </w:tblBorders>
        ''')
        tblPr[0].append(borders)

def insert_image_block(doc, img_filename, caption_text):
    img_path = os.path.join(PNG_DIR, img_filename)
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(6)
        p_img.paragraph_format.space_after = Pt(3)
        p_img.paragraph_format.keep_with_next = True
        r_img = p_img.add_run()
        r_img.add_picture(img_path, width=Inches(5.8))
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(0)
        p_cap.paragraph_format.space_after = Pt(10)
        r_cap = p_cap.add_run(caption_text)
        r_cap.font.name = "Calibri"
        r_cap.font.size = Pt(9)
        r_cap.italic = True
        r_cap.font.color.rgb = RGBColor(71, 85, 105)

def build_docx():
    doc = Document()
    
    # Page setup - A4 Portrait with 2.54cm margins
    section = doc.sections[0]
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    
    # Separate cover page header/footer
    section.different_first_page_header_footer = True
    
    # Header for regular pages
    header = section.header
    p_hdr = header.paragraphs[0]
    p_hdr.text = "BUKU MANUAL PENGGUNA — SISTEM E-SCM MARMER TULUNGAGUNG"
    p_hdr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_hdr.runs[0].font.name = "Calibri"
    p_hdr.runs[0].font.size = Pt(8)
    p_hdr.runs[0].font.color.rgb = RGBColor(100, 116, 139)
    
    # Footer for regular pages
    footer = section.footer
    p_ftr = footer.paragraphs[0]
    p_ftr.text = "Hak Cipta © 2026 E-SCM Marmer Tulungagung | Luaran Penelitian & Dokumen HKI Kemenkumham RI"
    p_ftr.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_ftr.runs[0].font.name = "Calibri"
    p_ftr.runs[0].font.size = Pt(8)
    p_ftr.runs[0].font.color.rgb = RGBColor(148, 163, 184)
        
    # --- COVER PAGE ---
    cover_table = doc.add_table(rows=1, cols=1)
    cover_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_cell = cover_table.cell(0, 0)
    c_cell.width = Inches(6.5)
    set_cell_background(c_cell, "0F172A") # Dark navy slate
    set_cell_margins(c_cell, top=300, bottom=300, left=300, right=300)
    
    cp = c_cell.paragraphs[0]
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_before = Pt(20)
    cp.paragraph_format.space_after = Pt(10)
    
    # Logo
    logo_path = os.path.join(PNG_DIR, "logo.png")
    if os.path.exists(logo_path):
        run_logo = cp.add_run()
        run_logo.add_picture(logo_path, width=Inches(1.8))
        cp.add_run("\n\n")
        
    badge_run = cp.add_run("DOKUMEN PANDUAN RESMI PENGOPERASIAN SISTEM\n")
    badge_run.bold = True
    badge_run.font.name = "Calibri"
    badge_run.font.size = Pt(11)
    badge_run.font.color.rgb = RGBColor(56, 189, 248) # Sky blue
    
    title_run = cp.add_run("BUKU MANUAL PENGGUNA (USER GUIDE)\n")
    title_run.bold = True
    title_run.font.name = "Calibri"
    title_run.font.size = Pt(24)
    title_run.font.color.rgb = RGBColor(255, 255, 255)
    
    sub_run = cp.add_run("SISTEM E-SUPPLY CHAIN MANAGEMENT (E-SCM) &\nE-COMMERCE IKM MARMER & ONYX TULUNGAGUNG\n\n")
    sub_run.bold = True
    sub_run.font.name = "Calibri"
    sub_run.font.size = Pt(14)
    sub_run.font.color.rgb = RGBColor(226, 232, 240)
    
    desc_run = cp.add_run(
        "Panduan Komprehensif Seluruh Alur Operasional: Transaksi E-Commerce Midtrans Snap QRIS, "
        "Verifikasi 2-Gate SPK Produksi, Penjadwalan Papan Kanban, Pemantauan WIP, "
        "Inspeksi Mutu Quality Control 2-Tahap, Pengendalian Residu Limbah, "
        "Checklist Packing Peti Kayu Solid Distribusi, dan Peramalan Permintaan AI ARIMA(2,0,2).\n\n"
    )
    desc_run.font.name = "Calibri"
    desc_run.font.size = Pt(10)
    desc_run.font.color.rgb = RGBColor(148, 163, 184)
    
    divider = cp.add_run("────────────────────────────────────────────\n\n")
    divider.font.name = "Calibri"
    divider.font.color.rgb = RGBColor(71, 85, 105)
    
    meta_run = cp.add_run(
        "Mitra Empiris: UD Cahaya Onix & UD Putra Abadi\n"
        "Lokasi Riset: Sentra Industri Kerajinan Marmer Campurdarat, Tulungagung\n"
        "Versi Sistem: 1.2.0 (Updated Official Release)\n"
        "Lampiran Luaran Penelitian & Pendaftaran Hak Cipta (HKI) DJKI Kemenkumham RI\n"
        "Tahun 2026"
    )
    meta_run.bold = True
    meta_run.font.name = "Calibri"
    meta_run.font.size = Pt(9.5)
    meta_run.font.color.rgb = RGBColor(203, 213, 225)
    
    doc.add_page_break()
    
    # Read Markdown content
    with open(MD_PATH, 'r', encoding='utf-8') as f:
        md_content = f.read()
        
    lines = md_content.split('\n')
    
    in_table = False
    table_lines = []
    
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        
        # Check if table block
        if stripped.startswith('|') and stripped.endswith('|'):
            in_table = True
            table_lines.append(stripped)
            i += 1
            continue
        elif in_table:
            # Process table
            if len(table_lines) >= 2:
                # Parse markdown table
                headers = [c.strip() for c in table_lines[0].split('|')[1:-1]]
                rows_data = []
                for t_line in table_lines[2:]:
                    cells = [c.strip() for c in t_line.split('|')[1:-1]]
                    rows_data.append(cells)
                    
                table = doc.add_table(rows=len(rows_data)+1, cols=len(headers))
                for col_idx, h in enumerate(headers):
                    table.cell(0, col_idx).text = h
                for row_idx, r_data in enumerate(rows_data):
                    for col_idx, val in enumerate(r_data):
                        if col_idx < len(headers):
                            table.cell(row_idx+1, col_idx).text = val
                format_table(table)
                p_spc = doc.add_paragraph()
                p_spc.paragraph_format.space_before = Pt(2)
                p_spc.paragraph_format.space_after = Pt(6)
            in_table = False
            table_lines = []
            
        # Check for headings
        if stripped.startswith('# '):
            i += 1
            continue
        elif stripped.startswith('## '):
            heading_text = stripped[3:].strip()
            h = doc.add_heading(level=1)
            h.paragraph_format.space_before = Pt(18)
            h.paragraph_format.space_after = Pt(8)
            h.paragraph_format.keep_with_next = True
            run = h.add_run(heading_text)
            run.font.name = "Calibri"
            run.font.size = Pt(16)
            run.bold = True
            run.font.color.rgb = RGBColor(30, 58, 138) # Navy Blue
            i += 1
            continue
        elif stripped.startswith('### '):
            heading_text = stripped[4:].strip()
            h = doc.add_heading(level=2)
            h.paragraph_format.space_before = Pt(12)
            h.paragraph_format.space_after = Pt(4)
            h.paragraph_format.keep_with_next = True
            run = h.add_run(heading_text)
            run.font.name = "Calibri"
            run.font.size = Pt(13)
            run.bold = True
            run.font.color.rgb = RGBColor(15, 23, 42) # Slate Dark
            
            # Check if any screenshot matches this subsection
            for mapping in SCREENSHOT_MAPPINGS:
                if re.search(mapping["pattern"], stripped):
                    insert_image_block(doc, mapping["image"], mapping["caption"])
                        
            i += 1
            continue
        elif stripped.startswith('#### '):
            heading_text = stripped[5:].strip()
            h = doc.add_heading(level=3)
            h.paragraph_format.space_before = Pt(8)
            h.paragraph_format.space_after = Pt(2)
            h.paragraph_format.keep_with_next = True
            run = h.add_run(heading_text)
            run.font.name = "Calibri"
            run.font.size = Pt(11)
            run.bold = True
            run.font.color.rgb = RGBColor(51, 65, 85)
            i += 1
            continue
            
        # Check for inline screenshot triggers in non-heading lines (e.g. Kartu Produk, Midtrans Snap)
        for mapping in SCREENSHOT_MAPPINGS:
            if not mapping["pattern"].startswith("###") and re.search(mapping["pattern"], stripped):
                insert_image_block(doc, mapping["image"], mapping["caption"])

        # Check for ascii art / code blocks
        if stripped.startswith('```'):
            code_lines = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith('```'):
                code_lines.append(lines[i])
                i += 1
            code_text = "\n".join(code_lines)
            
            p_code = doc.add_paragraph()
            p_code.paragraph_format.space_before = Pt(4)
            p_code.paragraph_format.space_after = Pt(6)
            run_code = p_code.add_run(code_text)
            run_code.font.name = "Consolas"
            run_code.font.size = Pt(8.5)
            run_code.font.color.rgb = RGBColor(30, 41, 59)
            i += 1
            continue
            
        # Check for bullet points
        if stripped.startswith('- ') or stripped.startswith('* '):
            bullet_text = stripped[2:].strip()
            p_bullet = doc.add_paragraph(style='List Bullet')
            p_bullet.paragraph_format.space_before = Pt(1)
            p_bullet.paragraph_format.space_after = Pt(2)
            
            parts = re.split(r'(\*\*.*?\*\*)', bullet_text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    r = p_bullet.add_run(part[2:-2])
                    r.bold = True
                else:
                    r = p_bullet.add_run(part)
                r.font.name = "Calibri"
                r.font.size = Pt(10)
                r.font.color.rgb = RGBColor(30, 41, 59)
            i += 1
            continue
            
        # Check for numbered list
        num_match = re.match(r'^(\d+)\.\s+(.*)$', stripped)
        if num_match:
            num = num_match.group(1)
            num_text = num_match.group(2)
            p_num = doc.add_paragraph(style='List Number')
            p_num.paragraph_format.space_before = Pt(1)
            p_num.paragraph_format.space_after = Pt(2)
            
            parts = re.split(r'(\*\*.*?\*\*)', num_text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    r = p_num.add_run(part[2:-2])
                    r.bold = True
                else:
                    r = p_num.add_run(part)
                r.font.name = "Calibri"
                r.font.size = Pt(10)
                r.font.color.rgb = RGBColor(30, 41, 59)
            i += 1
            continue
            
        # Regular paragraph
        if stripped:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.15
            
            parts = re.split(r'(\*\*.*?\*\*|\*.*?\*)', stripped)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    r = p.add_run(part[2:-2])
                    r.bold = True
                elif part.startswith('*') and part.endswith('*'):
                    r = p.add_run(part[1:-1])
                    r.italic = True
                else:
                    r = p.add_run(part)
                r.font.name = "Calibri"
                r.font.size = Pt(10)
                r.font.color.rgb = RGBColor(30, 41, 59)
                
        i += 1
        
    os.makedirs(os.path.dirname(DOCX_OUTPUT), exist_ok=True)
    doc.save(DOCX_OUTPUT)
    print(f"SUCCESS: Buku Manual Pengguna DOCX created successfully: {DOCX_OUTPUT}")

if __name__ == "__main__":
    build_docx()
